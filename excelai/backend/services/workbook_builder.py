"""Deterministic renderer: WorkbookSpec -> polished .xlsx (openpyxl only).

This is the engine that gives ExcelLence "Claude-style" Excel output — multi-sheet
workbooks with styled headers, number formats, banded rows, frozen panes, autofilters,
native charts, conditional formatting, and an executive dashboard sheet.

No new dependencies: charts/formatting all ship with openpyxl.
"""
from __future__ import annotations

import io
import re
from typing import Any

from openpyxl import Workbook
from openpyxl.chart import AreaChart, BarChart, LineChart, PieChart, Reference
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from backend.models.workbook_schemas import (
    ChartSpec,
    ColumnSpec,
    ConditionalRule,
    DashboardSpec,
    SheetSpec,
    WorkbookSpec,
)

THEMES = {
    "corporate": {"header_fill": "1F4E78", "header_font": "FFFFFF", "accent": "2E75B6", "band": "EAF1F8"},
    "dark": {"header_fill": "0D0F14", "header_font": "FFFFFF", "accent": "00C853", "band": "F2F5F7"},
    "light": {"header_fill": "EFEFEF", "header_font": "1A1D23", "accent": "4472C4", "band": "F9FAFB"},
}

NUMBER_FORMATS = {
    "currency": "$#,##0.00",
    "percentage": "0.00%",
    "date": "yyyy-mm-dd",
    "number": "#,##0.00",
    "text": None,
}

_CHART_CLASSES = {"bar": BarChart, "line": LineChart, "pie": PieChart, "area": AreaChart}


# ---- shared helpers ---------------------------------------------------------

def to_number(value: Any) -> float | None:
    """Best-effort numeric coercion. '$1,234' -> 1234.0, '12.5%' -> 0.125."""
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if not text:
        return None
    is_pct = text.endswith("%")
    cleaned = re.sub(r"[^0-9.\-]", "", text)
    if cleaned in ("", "-", ".", "-."):
        return None
    try:
        number = float(cleaned)
    except ValueError:
        return None
    return number / 100 if is_pct else number


def _coerce(value: Any, column_type: str) -> Any:
    if column_type in ("number", "currency", "percentage"):
        number = to_number(value)
        return number if number is not None else value
    return value


def _safe_title(name: str) -> str:
    cleaned = re.sub(r"[:\\/?*\[\]]", " ", str(name)).strip() or "Sheet"
    return cleaned[:31]


# ---- data sheet -------------------------------------------------------------

def _style_header(ws, headers: list[str], theme: dict) -> None:
    fill = PatternFill("solid", fgColor=theme["header_fill"])
    font = Font(bold=True, color=theme["header_font"])
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 20


def _apply_column_formats(ws, columns: list[ColumnSpec], n_rows: int, theme: dict) -> None:
    band = PatternFill("solid", fgColor=theme["band"])
    for r in range(2, n_rows + 2):
        for c, column in enumerate(columns, start=1):
            cell = ws.cell(row=r, column=c)
            fmt = NUMBER_FORMATS.get(column.type)
            if fmt:
                cell.number_format = fmt
            if column.type in ("number", "currency", "percentage"):
                cell.alignment = Alignment(horizontal="right")
            if r % 2 == 0:
                cell.fill = band


def _autosize(ws) -> None:
    widths: dict[str, int] = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None or not hasattr(cell, "column_letter"):
                continue
            length = len(str(cell.value))
            widths[cell.column_letter] = min(max(widths.get(cell.column_letter, 0), length), 48)
    for letter, width in widths.items():
        ws.column_dimensions[letter].width = width + 2


def _apply_conditional(ws, rule: ConditionalRule, headers: list[str], n_rows: int) -> None:
    if rule.column not in headers or n_rows < 1:
        return
    col_letter = get_column_letter(headers.index(rule.column) + 1)
    cell_range = f"{col_letter}2:{col_letter}{n_rows + 1}"
    if rule.rule == "color_scale":
        ws.conditional_formatting.add(cell_range, ColorScaleRule(
            start_type="min", start_color="F8696B",
            mid_type="percentile", mid_value=50, mid_color="FFEB84",
            end_type="max", end_color="63BE7B",
        ))
    elif rule.rule == "data_bar":
        ws.conditional_formatting.add(cell_range, DataBarRule(
            start_type="min", end_type="max", color=rule.color or "2E75B6",
        ))
    elif rule.rule in ("greater_than", "less_than") and rule.value is not None:
        operator = "greaterThan" if rule.rule == "greater_than" else "lessThan"
        ws.conditional_formatting.add(cell_range, CellIsRule(
            operator=operator, formula=[str(rule.value)],
            fill=PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
        ))


def _add_chart(target_ws, data_ws, spec: ChartSpec, headers: list[str], n_rows: int, anchor: str) -> None:
    if n_rows < 1 or spec.category_column not in headers:
        return
    chart_cls = _CHART_CLASSES.get(spec.type, BarChart)
    chart = chart_cls()
    chart.title = spec.title
    chart.height = 8
    chart.width = 16

    last_row = n_rows + 1
    cat_idx = headers.index(spec.category_column) + 1
    categories = Reference(data_ws, min_col=cat_idx, min_row=2, max_row=last_row)

    value_columns = spec.value_columns[:1] if spec.type == "pie" else spec.value_columns
    added = False
    for value_column in value_columns:
        if value_column not in headers:
            continue
        v_idx = headers.index(value_column) + 1
        data = Reference(data_ws, min_col=v_idx, min_row=1, max_row=last_row)
        chart.add_data(data, titles_from_data=True)
        added = True
    if not added:
        return
    chart.set_categories(categories)
    target_ws.add_chart(chart, anchor)


def _render_data_sheet(wb: Workbook, sheet: SheetSpec, theme: dict):
    ws = wb.create_sheet(title=_safe_title(sheet.name))
    columns = sheet.columns or [ColumnSpec(name=f"Column {i + 1}") for i in range(len(sheet.rows[0]) if sheet.rows else 1)]
    headers = [c.name for c in columns]
    ws.append(headers)

    for row in sheet.rows:
        ws.append([_coerce(row[i], columns[i].type) if i < len(columns) else row[i] for i in range(len(row))])

    n_rows = len(sheet.rows)
    _style_header(ws, headers, theme)
    _apply_column_formats(ws, columns, n_rows, theme)

    for cell_ref, formula in (sheet.formulas or {}).items():
        try:
            ws[cell_ref] = formula
        except Exception:
            continue

    if sheet.freeze_header:
        ws.freeze_panes = "A2"
    if sheet.autofilter and headers:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, n_rows + 1)}"

    for rule in sheet.conditional_formats:
        _apply_conditional(ws, rule, headers, n_rows)

    _autosize(ws)

    anchor_col = get_column_letter(len(headers) + 2)
    for i, chart in enumerate(sheet.charts):
        _add_chart(ws, ws, chart, headers, n_rows, chart.anchor or f"{anchor_col}{2 + i * 16}")

    return ws, headers, n_rows


# ---- dashboard sheet --------------------------------------------------------

def _render_dashboard(wb: Workbook, dash: DashboardSpec, data_sheets: dict, theme: dict):
    ws = wb.create_sheet(title=_safe_title(dash.name or "Dashboard"))
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:H1")
    headline = ws["A1"]
    headline.value = dash.headline or "Executive Dashboard"
    headline.font = Font(bold=True, size=18, color="1A1D23")
    headline.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 30

    card_fill = PatternFill("solid", fgColor=theme["band"])
    col = 1
    kpi_row = 3
    for kpi in dash.kpis[:8]:
        for rr in range(kpi_row, kpi_row + 3):
            for cc in range(col, col + 2):
                ws.cell(row=rr, column=cc).fill = card_fill
        ws.cell(row=kpi_row, column=col, value=kpi.label).font = Font(size=9, bold=True, color="6B7280")
        ws.cell(row=kpi_row + 1, column=col, value=kpi.value).font = Font(size=16, bold=True, color="1A1D23")
        if kpi.delta:
            up = str(kpi.delta).strip().startswith("+")
            ws.cell(row=kpi_row + 2, column=col, value=kpi.delta).font = Font(
                size=9, bold=True, color="1F9D55" if up else "E02424"
            )
        col += 3
        if col > 10:
            col = 1
            kpi_row += 4

    insights_row = kpi_row + 4
    ws.cell(row=insights_row, column=1, value="Key Insights").font = Font(bold=True, size=12, color="1A1D23")
    insights_row += 1
    for line in dash.narrative[:12]:
        ws.merge_cells(start_row=insights_row, start_column=1, end_row=insights_row, end_column=8)
        ws.cell(row=insights_row, column=1, value=f"•  {line}").font = Font(size=10, color="374151")
        insights_row += 1

    source = dash.source_sheet or (next(iter(data_sheets)) if data_sheets else None)
    if source and source in data_sheets:
        data_ws, headers, n_rows = data_sheets[source]
        chart_row = insights_row + 2
        for i, chart in enumerate(dash.charts):
            _add_chart(ws, data_ws, chart, headers, n_rows, chart.anchor or f"A{chart_row + i * 16}")

    return ws


# ---- entry point ------------------------------------------------------------

def build_workbook(spec: WorkbookSpec, filename: str | None = None) -> tuple[bytes, str]:
    theme = THEMES.get(spec.theme, THEMES["corporate"])
    wb = Workbook()
    wb.remove(wb.active)  # drop the default empty sheet

    data_sheets: dict[str, tuple] = {}
    for sheet in spec.sheets:
        ws, headers, n_rows = _render_data_sheet(wb, sheet, theme)
        data_sheets[sheet.name] = (ws, headers, n_rows)

    if spec.dashboard:
        _render_dashboard(wb, spec.dashboard, data_sheets, theme)
        # surface the dashboard as the first tab
        dash_title = _safe_title(spec.dashboard.name or "Dashboard")
        wb._sheets.sort(key=lambda s: 0 if s.title == dash_title else 1)

    if not wb.sheetnames:
        wb.create_sheet("Sheet1")

    buffer = io.BytesIO()
    wb.save(buffer)

    name = filename or f"{spec.title or 'ExcelLence'}.xlsx"
    if not name.lower().endswith(".xlsx"):
        name += ".xlsx"
    return buffer.getvalue(), name
